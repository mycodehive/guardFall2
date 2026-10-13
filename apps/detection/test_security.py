"""Boundary tests for private monitoring data and authenticated mutations."""
import json
import tempfile
from datetime import timedelta

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group
from cryptography.fernet import Fernet
from django.core.files.base import ContentFile
from django.test import Client, TestCase, override_settings
from django.utils import timezone

from apps.menus.models import Menu
from .ml import JOINTS
from .models import DetectionSettings, FallEvent, Job, MonitoringSession, NotificationSettings, PoseFrame


class DetectionSecurityTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        users = get_user_model()
        cls.owner = users.objects.create_user('owner', password='Testing-private-12')
        cls.other = users.objects.create_user('other', password='Testing-private-12')
        cls.staff = users.objects.create_user('staff', password='Testing-private-12', is_staff=True)
        for code in ('monitor', 'events', 'jobs', 'models', 'settings', 'notifications'):
            Menu.objects.update_or_create(code=code, defaults={'title': code, 'url': '/', 'active': True})

    def setUp(self):
        self.client.force_login(self.owner)
        self.session = MonitoringSession.objects.create(owner=self.owner, config={
            'MIN_OK_PARTS': 3, 'FALL_COUNT': 3, 'LS_Y': .55, 'RS_Y': .55,
            'LK_Y': .65, 'RK_Y': .65, 'VISIBILITY': .7,
            'FALL_PROBABILITY': .5, 'consecutive_frames': 2, 'cooldown_seconds': 30,
        })

    def joints(self, y=.9, visible=1):
        return {f'{joint}_{axis}': value for joint in JOINTS
                for axis, value in (('x', .5), ('y', y), ('v', visible), ('vr', visible))}

    def frame(self, sequence, joints=None):
        return self.client.post(f'/monitor/{self.session.pk}/frame/',
                                data=json.dumps({'sequence': sequence, 'joints': joints or self.joints()}),
                                content_type='application/json')

    def test_anonymous_monitor_redirects_to_login(self):
        response = Client().get('/monitor/')
        self.assertEqual(response.status_code, 302)
        self.assertIn('/accounts/login/', response['Location'])

    def test_regular_user_cannot_open_staff_settings(self):
        self.assertEqual(self.client.get('/settings/').status_code, 403)

    def test_other_user_cannot_submit_or_stop_private_session(self):
        self.client.force_login(self.other)
        self.assertEqual(self.frame(0).status_code, 404)
        self.assertEqual(self.client.post(f'/monitor/{self.session.pk}/stop/').status_code, 404)
        self.session.refresh_from_db()
        self.assertTrue(self.session.active)
        self.assertFalse(PoseFrame.objects.filter(session=self.session).exists())

    def test_private_job_download_denies_other_user(self):
        with tempfile.TemporaryDirectory() as directory, override_settings(MEDIA_ROOT=directory):
            job = Job.objects.create(owner=self.staff, kind='video', name='private', status='done')
            job.result_file.save('result.csv', ContentFile(b'label\n1\n'))
            other_staff = get_user_model().objects.create_user('otherstaff', is_staff=True)
            self.client.force_login(other_staff)
            self.assertEqual(self.client.get(f'/jobs/{job.pk}/download/').status_code, 404)
            self.client.force_login(self.staff)
            response = self.client.get(f'/jobs/{job.pk}/download/')
            self.assertEqual(response.status_code, 200)
            self.assertEqual(b''.join(response.streaming_content), b'label\n1\n')
            response.close()

    def test_other_user_cannot_review_private_event(self):
        frame = PoseFrame.objects.create(session=self.session, sequence=0, joints=self.joints(), prediction=1)
        event = FallEvent.objects.create(session=self.session, frame=frame)
        self.client.force_login(self.other)
        response = self.client.post(f'/events/{event.pk}/review/', {'status': 'confirmed'})
        self.assertEqual(response.status_code, 404)
        event.refresh_from_db()
        self.assertEqual(event.status, 'new')

    def test_mutations_require_csrf_token(self):
        client = Client(enforce_csrf_checks=True)
        client.force_login(self.owner)
        for path in ('/monitor/start/', f'/monitor/{self.session.pk}/frame/', f'/monitor/{self.session.pk}/stop/'):
            with self.subTest(path=path):
                self.assertEqual(client.post(path, '{}', content_type='application/json').status_code, 403)

    def test_invalid_sequence_or_missing_joints_rejected_without_persistence(self):
        payloads = [ {'sequence': -1, 'joints': self.joints()},
                     {'sequence': True, 'joints': self.joints()},
                     {'sequence': 1.5, 'joints': self.joints()},
                     {'sequence': 0, 'joints': {}},
                     {'sequence': 0, 'joints': []} ]
        for payload in payloads:
            with self.subTest(payload=payload):
                response = self.client.post(f'/monitor/{self.session.pk}/frame/', json.dumps(payload), content_type='application/json')
                self.assertEqual(response.status_code, 400)
        self.assertFalse(self.session.frames.exists())

    def test_nonfinite_joint_values_rejected(self):
        for invalid in (float('nan'), float('inf'), -float('inf')):
            joints = self.joints()
            joints['left_shoulder_y'] = invalid
            with self.subTest(invalid=invalid):
                self.assertEqual(self.frame(0, joints).status_code, 400)
        self.assertFalse(self.session.frames.exists())

    def test_replayed_or_out_of_order_frame_cannot_change_state(self):
        self.assertEqual(self.frame(2).status_code, 200)
        for sequence in (2, 1):
            self.assertEqual(self.frame(sequence).status_code, 409)
        self.session.refresh_from_db()
        self.assertEqual(self.session.last_sequence, 2)
        self.assertEqual(self.session.fall_streak, 1)
        self.assertEqual(self.session.frames.count(), 1)

    def test_streak_and_cooldown_gate_events(self):
        for sequence in range(4):
            self.assertEqual(self.frame(sequence).status_code, 200)
            self.assertEqual(self.session.events.count(), int(sequence >= 1))
        self.session.refresh_from_db()
        self.session.last_event_at = timezone.now() - timedelta(seconds=31)
        self.session.save(update_fields=['last_event_at'])
        self.assertEqual(self.frame(4).status_code, 200)
        self.assertEqual(self.session.events.count(), 2)

    def test_normal_and_unknown_frames_reset_fall_streak(self):
        self.assertEqual(self.frame(0).status_code, 200)
        self.assertEqual(self.frame(1, self.joints(y=.1)).status_code, 200)
        self.session.refresh_from_db()
        self.assertEqual(self.session.fall_streak, 0)
        self.assertEqual(self.frame(2).status_code, 200)
        self.assertEqual(self.frame(3, self.joints(visible=0)).status_code, 200)
        self.session.refresh_from_db()
        self.assertEqual(self.session.fall_streak, 0)
        self.assertFalse(self.session.events.exists())

    def test_stopped_session_rejects_frames(self):
        self.assertEqual(self.client.post(f'/monitor/{self.session.pk}/stop/').status_code, 200)
        self.assertEqual(self.frame(0).status_code, 409)
        self.session.refresh_from_db()
        self.assertFalse(self.session.active)
        self.assertIsNotNone(self.session.ended_at)

    def test_menu_ancestor_group_restriction_applies_to_direct_url(self):
        group = Group.objects.create(name='Restricted operators')
        parent = Menu.objects.create(code='restricted-parent', title='Private')
        parent.groups.add(group)
        menu = Menu.objects.get(code='monitor')
        menu.parent = parent
        menu.save(update_fields=['parent'])
        self.assertEqual(self.client.get('/monitor/').status_code, 403)
        self.owner.groups.add(group)
        self.assertEqual(self.client.get('/monitor/').status_code, 200)
        parent.active = False
        parent.save(update_fields=['active'])
        self.assertEqual(self.client.get('/monitor/').status_code, 403)

    def test_invalid_settings_do_not_modify_saved_configuration(self):
        cfg = DetectionSettings.objects.create(pk=1)
        self.client.force_login(self.staff)
        valid = {f.name: getattr(cfg, f.name) for f in cfg._meta.fields if f.name != 'id'}
        for invalid in ('nan', 'inf', '-1', '1.1'):
            with self.subTest(invalid=invalid):
                response = self.client.post('/settings/', {**valid, 'visibility': invalid})
                self.assertEqual(response.status_code, 200)
                self.assertTrue(response.context['form'].errors)
                cfg.refresh_from_db()
                self.assertEqual(cfg.visibility, .7)

    def test_notification_secrets_are_encrypted_and_not_rendered_or_audited(self):
        from apps.core.models import AuditLog
        from .notifications import decrypt
        token = '123456:secret_test_token'
        chat = 'secret_test_chat'
        with override_settings(GUARDFALL_ENCRYPTION_KEY=Fernet.generate_key().decode()):
            self.assertEqual(self.client.post('/notifications/', {'enabled': 'on', 'token': token, 'chat_id': chat}).status_code, 302)
            cfg = NotificationSettings.objects.get(owner=self.owner)
            self.assertNotEqual(cfg.token_ciphertext, token)
            self.assertNotEqual(cfg.chat_id_ciphertext, chat)
            self.assertEqual(decrypt(cfg.token_ciphertext), token)
            self.assertEqual(decrypt(cfg.chat_id_ciphertext), chat)
            response = self.client.get('/notifications/')
            self.assertNotContains(response, token)
            self.assertNotContains(response, chat)
            for record in AuditLog.objects.filter(action='notifications.update').values():
                self.assertNotIn(token, str(record))
                self.assertNotIn(chat, str(record))

    def test_spoofed_visibility_flag_is_ignored(self):
        joints = self.joints(visible=0)
        for joint in JOINTS:
            joints[f'{joint}_vr'] = 1
        self.assertEqual(self.frame(0, joints).status_code, 200)
        frame = self.session.frames.get()
        self.assertEqual(frame.prediction, -1)

    def test_formula_like_notes_are_escaped_in_excel_export(self):
        from io import BytesIO
        from openpyxl import load_workbook
        frame = PoseFrame.objects.create(session=self.session, sequence=0, joints=self.joints(), prediction=1)
        FallEvent.objects.create(session=self.session, frame=frame, note='=HYPERLINK("https://example.invalid")')
        response = self.client.get('/events/export/')
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.content))
        self.assertTrue(workbook.active.cell(2, 5).value.startswith("'=HYPERLINK"))
        self.assertEqual(workbook.active.cell(2, 5).data_type, 's')

    def test_camera_image_ownership_and_invalid_payload(self):
        import base64
        self.client.force_login(self.other)
        response = self.client.post(f'/monitor/{self.session.pk}/image/', '{}', content_type='application/json')
        self.assertEqual(response.status_code, 404)
        self.client.force_login(self.owner)
        for image in ('%%%not-base64%%%', base64.b64encode(b'not-an-image').decode()):
            with self.subTest(image=image):
                response = self.client.post(f'/monitor/{self.session.pk}/image/', json.dumps({'sequence': 0, 'image': image}), content_type='application/json')
                self.assertEqual(response.status_code, 400)
        self.assertFalse(self.session.frames.exists())

    def test_camera_image_uses_same_ingestion_and_does_not_store_image(self):
        import base64
        from unittest.mock import patch
        raw = b'example-image-bytes'
        with patch('apps.detection.ml.extract_image', return_value=self.joints()) as extract:
            response = self.client.post(f'/monitor/{self.session.pk}/image/', json.dumps({'sequence': 0, 'image': base64.b64encode(raw).decode()}), content_type='application/json')
        self.assertEqual(response.status_code, 200)
        extract.assert_called_once_with(raw, self.session.config)
        frame = self.session.frames.get()
        self.assertEqual(frame.prediction, 1)
        self.assertNotIn('image', frame.joints)
        self.assertEqual(self.frame(0).status_code, 409)

    def test_expired_session_rejects_frame_without_persistence(self):
        MonitoringSession.objects.filter(pk=self.session.pk).update(started_at=timezone.now()-timedelta(hours=9))
        self.assertEqual(self.frame(0).status_code, 409)
        self.assertFalse(self.session.frames.exists())

    def test_seeded_navigation_has_three_levels(self):
        monitor = Menu.objects.get(code='monitor')
        self.assertEqual(monitor.parent.code, 'operations')
        self.assertEqual(monitor.parent.parent.code, 'safety')
        self.assertIsNone(monitor.parent.parent.parent)
        self.assertEqual(Menu.objects.get(code='settings').parent.code, 'config')

    def test_model_selection_requires_owner_or_explicit_assignment(self):
        from .models import TrainedModel
        model = TrainedModel.objects.create(owner=self.staff, name='Assigned classifier', algorithm='dense', active=True, artifact='private/model.zip')
        before = MonitoringSession.objects.filter(owner=self.owner).count()
        self.assertEqual(self.client.post('/monitor/start/', {'model': model.pk}).status_code, 404)
        self.assertEqual(MonitoringSession.objects.filter(owner=self.owner).count(), before)
        model.allowed_users.add(self.owner)
        response = self.client.post('/monitor/start/', {'model': model.pk})
        self.assertEqual(response.status_code, 200)
        session = MonitoringSession.objects.get(pk=response.json()['id'])
        self.assertEqual(session.owner, self.owner)
        self.assertEqual(session.model, model)
        model.active = False
        model.save(update_fields=['active'])
        self.assertEqual(self.client.post('/monitor/start/', {'model': model.pk}).status_code, 404)

    def test_staff_and_regular_user_page_templates(self):
        for path in ('/', '/monitor/', '/events/', '/notifications/', '/accounts/password/'):
            with self.subTest(role='user', path=path):
                self.assertEqual(self.client.get(path).status_code, 200)
        for path in ('/jobs/', '/models/', '/settings/'):
            with self.subTest(role='user denied', path=path):
                self.assertEqual(self.client.get(path).status_code, 403)
        self.client.force_login(self.staff)
        for path in ('/', '/monitor/', '/events/', '/notifications/', '/accounts/password/', '/jobs/', '/jobs/?kind=train', '/models/', '/settings/'):
            with self.subTest(role='staff', path=path):
                self.assertEqual(self.client.get(path).status_code, 200)

    def test_menu_cycle_and_fourth_level_are_invalid(self):
        from django.core.exceptions import ValidationError
        root = Menu.objects.create(code='depth-root', title='Root')
        middle = Menu.objects.create(code='depth-middle', title='Middle', parent=root)
        leaf = Menu.objects.create(code='depth-leaf', title='Leaf', parent=middle)
        fourth = Menu(code='depth-fourth', title='Fourth', parent=leaf)
        with self.assertRaises(ValidationError):
            fourth.full_clean()
        root.parent = leaf
        with self.assertRaises(ValidationError):
            root.full_clean()

    def test_dashboard_hides_events_when_menu_disabled(self):
        frame = PoseFrame.objects.create(session=self.session, sequence=0, joints=self.joints(), prediction=1)
        FallEvent.objects.create(session=self.session, frame=frame)
        menu = Menu.objects.get(code='events')
        menu.active = False
        menu.save(update_fields=['active'])
        response = self.client.get('/')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context['event_count'], 0)
        self.assertEqual(list(response.context['recent_events']), [])
